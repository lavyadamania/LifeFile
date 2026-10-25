from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.enum.section import WD_ORIENT

OUTPUT = 'LIFEFILE_Unified_Project_Proposal_Reference_Format.docx'

NAVY = '17365D'
BLUE = '1F4E79'
TEAL = '0F6B78'
LIGHT_BLUE = 'D9EAF7'
LIGHT_TEAL = 'DDEBF7'
LIGHT_GREY = 'F2F4F7'
MID_GREY = '667085'
DARK = '1F2937'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in kwargs[edge]:
                    element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))


def set_cell_text(cell, text, bold=False, color=DARK, size=9, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        set_cell_text(header.cells[idx], value, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[idx], NAVY)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            set_cell_text(row.cells[idx], value, size=font_size)
            if len(table.rows) % 2 == 0:
                set_cell_shading(row.cells[idx], LIGHT_GREY)
    if widths:
        set_table_widths(table, widths)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={'val': 'single', 'sz': '4', 'color': 'D0D5DD'}, bottom={'val': 'single', 'sz': '4', 'color': 'D0D5DD'}, left={'val': 'single', 'sz': '4', 'color': 'D0D5DD'}, right={'val': 'single', 'sz': '4', 'color': 'D0D5DD'})
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = None
    p.paragraph_format.line_spacing = None
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_callout(doc, title, text, fill=LIGHT_TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, left={'val': 'single', 'sz': '18', 'color': TEAL}, top={'val': 'single', 'sz': '4', 'color': 'B7C9D6'}, bottom={'val': 'single', 'sz': '4', 'color': 'B7C9D6'}, right={'val': 'single', 'sz': '4', 'color': 'B7C9D6'})
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r.font.size = Pt(10)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.08
    for run in p2.runs:
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Page ')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)


def setup_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.25)
    sec.header_distance = Inches(0.5)
    sec.footer_distance = Inches(0.5)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = None
    normal.paragraph_format.line_spacing = None

    for name, size, color in [('Heading 1', 14, '365F91'), ('Heading 2', 13, '4F81BD'), ('Heading 3', 12, '4F81BD')]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(24 if name == 'Heading 1' else 10)
        style.paragraph_format.space_after = None
        style.paragraph_format.keep_with_next = True

    for style_name in ['List Bullet', 'List Bullet 2', 'List Number']:
        style = doc.styles[style_name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor.from_string(DARK)

    sec.header.paragraphs[0].text = ''
    sec.footer.paragraphs[0].text = ''


def title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run('LIFEFILE')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p2.add_run('Integrated Project Proposal\nPBL Framework, Department of Information Technology')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)


def contents_page(doc):
    add_heading(doc, 'Contents', 1)
    add_para(doc, 'The following structure mirrors a formal project proposal while retaining the technical depth required for architecture review, implementation planning, and live demonstration.')
    sections = [
        '1. Project Title', '2. Executive Summary', '3. Introduction and Context', '4. Problem Statement', '5. Motivation and Significance', '6. Objectives', '7. Scope of the Project', '8. Existing System and Gap Analysis', '9. Proposed System', '10. System Architecture', '11. Functional Modules', '12. Clinical Queue and ETA Methodology', '13. Patient Memory and Safety Workflow', '14. Security, Privacy, and Governance', '15. Data Model and API Surface', '16. Technology Stack and Tools', '17. Development Methodology', '18. Validation and Testing Strategy', '19. Feasibility Study', '20. Expected Outcomes', '21. Innovation and Novelty', '22. Limitations and Responsible-Use Boundaries', '23. Project Timeline', '24. Team Members and Responsibilities', '25. References', '26. Conclusion'
    ]
    for item in sections:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)
    doc.add_page_break()


def build():
    doc = Document()
    setup_document(doc)
    title_page(doc)

    add_heading(doc, '1. Project Title', 1)
    add_para(doc, 'LifeFile: An Intelligent, Real-Time Smart Clinic Operating System for Patient Flow, Clinical Context, and Multi-Facility Healthcare Operations')

    add_heading(doc, '2. Executive Summary', 1)
    add_para(doc, 'LifeFile, also referred to as SCOS (Smart Clinic Operating System), is a full-stack healthcare operations platform designed to coordinate patients, doctors, hospitals, clinics, and administrators through a single digital workflow. The platform addresses a practical operational problem: healthcare delivery is not only a medical activity, but also a complex coordination problem involving appointments, time-aware check-in, queue fairness, clinical urgency, patient records, communication, and institutional accountability.')
    add_para(doc, 'The implemented system combines a React and TypeScript frontend with a Node.js and Express backend, MongoDB persistence, JWT-based authentication, Kafka event streaming, Socket.IO real-time delivery, and a patient-memory subsystem that extracts structured context from clinical records. Its central operational capability is the Adaptive Clinical Priority Allocation (ACPA) queue engine, which separates a permanent appointment token from a dynamic service position. This permits urgent cases to receive appropriate priority while retaining fairness mechanisms such as wait-time aging and bounded missed-call penalties.')
    add_para(doc, 'LifeFile is deliberately positioned as an assistive clinical operations prototype rather than a replacement for a hospital information system or a medical professional. Triage, symptom analysis, risk estimation, and memory extraction are implemented through a combination of deterministic rules, local NLP, and optional generative-AI assistance. The platform therefore provides a strong foundation for demonstration and further research, while production deployment would require additional security hardening, clinical validation, compliance review, observability, and infrastructure controls.')
    add_callout(doc, 'Core proposition', 'LifeFile turns fragmented clinic operations into a coordinated, observable workflow: discover care -> book or check in -> prioritise safely -> serve in real time -> record the encounter -> preserve patient context -> audit the outcome.')

    add_heading(doc, '3. Introduction and Context', 1)
    add_para(doc, 'Healthcare facilities frequently experience operational friction even when qualified clinicians and appropriate infrastructure are available. Patients may arrive too early, queues may be managed through static tokens, urgent cases may not be surfaced quickly enough, and information from previous encounters may remain scattered across records. These problems affect patient experience, clinical safety, staff workload, and the ability of administrators to understand service performance.')
    add_para(doc, 'LifeFile responds to this environment by treating the clinic as a connected operational ecosystem. Its patient-facing experience supports discovery, booking, check-in, live queue visibility, records, medical timelines, symptom triage, risk education, and voice-assisted navigation. Its doctor-facing experience supports queue control, consultation documentation, prescription workflows, clinical memory review, and analytics. Hospital and administrator experiences provide facility management, roster administration, audit visibility, and comparative performance analysis.')
    add_para(doc, 'The project is especially relevant to settings in which multiple facilities share a healthcare network but must retain local operational boundaries. The system models hospitals, clinics, doctors, patients, appointments, records, prescriptions, memories, correction requests, reviews, join requests, and audit events as first-class entities. This provides a basis for extending the prototype toward a more comprehensive digital health operating model.')

    add_heading(doc, '4. Problem Statement', 1)
    add_para(doc, 'The project addresses five interconnected operational challenges:')
    add_number(doc, 'Static queue systems cannot express the difference between appointment order and clinical urgency. A patient with a high-risk complaint may remain behind routine cases simply because their token was issued later.')
    add_number(doc, 'Uncontrolled arrival and check-in behaviour creates unnecessary crowding. Patients need clear, time-aware information about when they may join a queue and what to expect.')
    add_number(doc, 'Patients often cannot see the current state of service delivery. Without a live position and bounded ETA, they must remain physically present or repeatedly contact staff.')
    add_number(doc, 'Clinical context is fragmented across prescriptions, notes, records, and visits. Important facts such as allergies or contradictory statements may be difficult to notice at the point of care.')
    add_number(doc, 'Facility operators require trustworthy administrative visibility. Without audit logs, role controls, analytics, and benchmark scenarios, it is difficult to evaluate performance or investigate actions.')
    add_para(doc, 'The central design problem is therefore to coordinate clinical urgency, fairness, patient communication, clinical context, and facility governance without presenting automated outputs as unquestionable clinical decisions.')

    add_heading(doc, '5. Motivation and Significance', 1)
    add_bullet(doc, 'Patient-centred care: provide transparent appointment, check-in, queue, and record experiences that reduce avoidable uncertainty.')
    add_bullet(doc, 'Operational efficiency: allow doctors and hospitals to work from a shared, authoritative state instead of disconnected lists and manual updates.')
    add_bullet(doc, 'Clinical safety: surface relevant patient context, including possible allergy conflicts and contradictory statements, before a prescription or consultation is completed.')
    add_bullet(doc, 'Fairness: combine clinical urgency with wait-time aging and bounded penalties so priority care does not become an excuse for indefinite routine-patient delay.')
    add_bullet(doc, 'Technical learning: demonstrate modern full-stack development, event-driven integration, applied NLP, data modelling, authentication, testing, and performance benchmarking in one coherent project.')
    add_bullet(doc, 'Research and evaluation value: compare queue policies such as FIFO, conventional priority, ACPA, and ablation variants under repeatable synthetic scenarios.')

    add_heading(doc, '6. Objectives', 1)
    add_heading(doc, '6.1 Primary Objectives', 2)
    add_bullet(doc, 'Develop a role-aware healthcare operations platform for patients, doctors, hospitals, and administrators.')
    add_bullet(doc, 'Implement appointment discovery, booking, rescheduling, cancellation, walk-in, and time-aware check-in workflows.')
    add_bullet(doc, 'Implement an ACPA queue engine that dynamically calculates service order while preserving permanent appointment tokens.')
    add_bullet(doc, 'Deliver live queue status, consultation state, and ETA changes across connected clients.')
    add_bullet(doc, 'Create a longitudinal patient-memory workflow with extraction, provenance, deduplication, contradiction flags, corrections, and review.')
    add_bullet(doc, 'Provide digital prescriptions, medical records, uploads, timelines, reviews, analytics, audit logs, and facility administration.')
    add_bullet(doc, 'Provide benchmark simulations and executable validation scenarios for queue and memory behaviour.')
    add_heading(doc, '6.2 Secondary Objectives', 2)
    add_bullet(doc, 'Maintain a clear separation between authoritative clinical records and derived patient-memory context.')
    add_bullet(doc, 'Keep automated assistance transparent, reviewable, and bounded by professional oversight.')
    add_bullet(doc, 'Design the implementation so that future hardening can address privacy, compliance, scale, observability, and deployment requirements.')

    add_heading(doc, '7. Scope of the Project', 1)
    add_heading(doc, '7.1 In Scope', 2)
    for item in ['Multi-role authentication and role-specific dashboards', 'Patient and hospital discovery, doctor profiles, booking, appointment management, and check-in', 'Dynamic doctor queues, emergency prioritisation, skipped-patient handling, consultation state, and ETA ranges', 'Real-time event delivery using KafkaJS and Socket.IO', 'Digital prescriptions, medical records, protected record access, uploads, and chronological timelines', 'Patient-memory extraction, source references, contradiction detection, correction requests, review, and audit events', 'Hospital, clinic, doctor, patient, roster, affiliation, review, and administrative workflows', 'Patient symptom triage, clinical-note parsing, risk education, and voice navigation assistance', 'Analytics, queue benchmarks, synthetic data seeding, database inspection, and acceptance tests']:
        add_bullet(doc, item)
    add_heading(doc, '7.2 Out of Scope', 2)
    for item in ['Autonomous diagnosis, prescribing, or treatment decisions', 'Replacement of a regulated electronic health record or hospital information system', 'Unsupervised clinical deployment or a claim of medical-model validation', 'Automatic merging, hiring, or other irreversible administrative decisions without human approval', 'Guarantees of zero security risk, zero queue delay, or universal interoperability with external health systems', 'Production-scale availability, disaster recovery, multi-region deployment, and formal regulatory certification']:
        add_bullet(doc, item)

    add_heading(doc, '8. Existing System and Gap Analysis', 1)
    add_para(doc, 'Many existing digital-health workflows address only one portion of the operational problem. Appointment portals help patients schedule visits but do not necessarily represent live queue state. Token systems make ordering visible but often remain FIFO and cannot respond to clinical urgency. Electronic records preserve documents but may not provide a concise, reviewable longitudinal context. Messaging tools can broadcast updates but do not create an authoritative workflow state.')
    add_table(doc, ['Observed approach', 'Strength', 'Operational gap addressed by LifeFile'], [
        ('Static FIFO token list', 'Simple and understandable', 'Cannot express urgency, wait aging, skipped calls, or dynamic service position'),
        ('Appointment-only portal', 'Supports scheduling', 'Does not coordinate check-in, live queue state, consultation progress, or ETA'),
        ('Document repository', 'Preserves files and records', 'Does not extract patient context or identify contradictory medical statements'),
        ('Standalone chatbot or symptom form', 'Offers an accessible interface', 'May be disconnected from booking, doctor workflow, audit trail, and human review'),
        ('Administrative spreadsheet workflow', 'Low initial setup cost', 'Limited role control, real-time synchronisation, traceability, and benchmarkability'),
    ], widths=[1.55, 2.0, 3.0])
    add_para(doc, 'LifeFile fills this gap by connecting operational state, patient-facing transparency, clinical context, and governance controls in one demonstrable system. The differentiator is not a single artificial-intelligence feature; it is the integration of multiple workflow stages into a coherent operating model.')

    add_heading(doc, '9. Proposed System', 1)
    add_para(doc, 'LifeFile is proposed as a layered, event-aware clinic operating system. The frontend presents role-specific workflows. The backend exposes REST APIs for authoritative commands and queries. MongoDB stores operational entities and clinical records. Kafka carries domain events, while Socket.IO distributes relevant state changes to connected clients. Supporting services implement queue scoring, ETA calculation, patient-memory extraction, and benchmark execution.')
    add_heading(doc, '9.1 End-to-End Operational Flow', 2)
    for item in ['A patient discovers a hospital or doctor and selects an available appointment slot.', 'The patient provides a chief complaint; the client-side classifier assigns an assistive triage level and stores the complaint with the appointment.', 'The patient checks in only within the configured time window. The backend records the authoritative status transition.', 'The doctor dashboard retrieves the active queue. ACPA calculates dynamic positions from token order, triage, waiting time, and missed calls.', 'A doctor calls or starts a consultation. The backend persists the state and publishes an event.', 'Kafka and Socket.IO deliver the change to relevant clients, which refresh or update their queue and ETA views.', 'The doctor records a prescription or consultation note. The patient-memory pipeline extracts possible structured facts and links them to source records.', 'Patients may view context and request corrections; authorised doctors or administrators review those requests.', 'Audit, analytics, and benchmark modules support oversight and continuous improvement.']:
        add_number(doc, item)
    add_heading(doc, '9.2 Role Model', 2)
    add_table(doc, ['Role', 'Primary responsibilities', 'Representative workspace'], [
        ('Patient', 'Discover care, book, check in, monitor queue, access records, review memory, request corrections', 'Patient dashboard, appointments, live queue, timeline, memory, analytics'),
        ('Doctor', 'Manage queue, call patients, conduct consultations, issue prescriptions, review context', 'Doctor dashboard, dynamic queue, consultation, schedule, memory panel'),
        ('Hospital / Clinic', 'Manage facility information, rosters, affiliations, appointments, and local analytics', 'Hospital dashboard, doctors, clinics, records, analytics'),
        ('Administrator', 'Manage network-wide entities, review audit activity, monitor statistics and benchmarks', 'Admin dashboard, hospitals, doctors, audit logs, benchmark, settings'),
    ], widths=[1.1, 3.55, 2.0])

    add_heading(doc, '10. System Architecture', 1)
    add_heading(doc, '10.1 Logical Architecture', 2)
    add_table(doc, ['Layer', 'Primary responsibility', 'Implemented technologies'], [
        ('Experience layer', 'Role-specific workflows, dashboards, forms, charts, notifications, accessibility-oriented controls', 'React 18, TypeScript, React Router, Tailwind CSS, Lucide, Recharts'),
        ('Client state and data layer', 'Authentication state, API calls, query caching, live event handling', 'Zustand, TanStack React Query, Axios, Socket.IO client'),
        ('Application layer', 'Authentication, appointment control, queue commands, records, memory, analytics, benchmarks', 'Node.js, Express, JWT middleware, route modules'),
        ('Domain services', 'Queue priority calculation, ETA estimation, Kafka connectivity, memory extraction', 'dynamicPriority.js, queueETA.js, kafka.js, memoryService.js'),
        ('Persistence layer', 'Operational and derived data models, uploads, audit information', 'MongoDB, Mongoose, local upload storage'),
        ('Event and integration layer', 'Durable domain-event transport and client-facing real-time broadcasts', 'KafkaJS, Apache Kafka, Socket.IO'),
    ], widths=[1.2, 3.25, 2.2])
    add_heading(doc, '10.2 Component Relationships', 2)
    add_para(doc, 'The frontend communicates with the backend through REST endpoints under `/api/*`. Commands that change state, such as check-in, call-next, skip, start consultation, and completion, are validated and persisted by the backend before corresponding events are published. Kafka consumers receive domain events and use Socket.IO to broadcast relevant updates. Clients then refetch authoritative state where appropriate, preventing the event stream from becoming an independent source of truth.')
    add_callout(doc, 'Architectural principle', 'MongoDB remains the source of truth for operational state. Kafka and Socket.IO communicate change and reduce perceived latency; they do not replace server-side authorisation or authoritative persistence.')
    add_heading(doc, '10.3 Deployment Topology for Demonstration', 2)
    add_table(doc, ['Process / service', 'Default responsibility', 'Typical local endpoint'], [
        ('Frontend development server', 'Serves the React application and development assets', 'http://localhost:5173'),
        ('Backend API server', 'Serves REST routes, Socket.IO, uploads, and health endpoint', 'http://localhost:5000'),
        ('MongoDB', 'Persists users, appointments, records, memories, and audit data', 'Configured through MONGO_URI'),
        ('Apache Kafka', 'Carries queue, appointment, prescription, and memory-related events', 'Configured through Kafka environment variables'),
    ], widths=[1.75, 3.55, 1.35])

    add_heading(doc, '11. Functional Modules', 1)
    modules = [
        ('11.1 Authentication and Role Control', 'Registration, login, JWT issuance, password hashing with bcryptjs, protected frontend routes, role-specific routing, and authenticated API access. Demo login shortcuts support presentation workflows but should not replace production identity management.'),
        ('11.2 Patient Discovery and Appointment Management', 'Hospital search, doctor profiles, available slots, appointment booking, rescheduling, cancellation, walk-ins, missed appointments, and status transitions. Chief complaints and triage metadata travel with the appointment record.'),
        ('11.3 Symptom Triage and Patient Assistance', 'A client-side clinical keyword classifier supports booking-time triage. The symptom checker uses compromise NLP and local rules for entity extraction and recommendations. The cardiovascular risk page provides educational risk estimation using a simplified Framingham-inspired model. Voice navigation uses browser speech capabilities and intent matching.'),
        ('11.4 Dynamic Queue Management', 'Doctors receive a live queue with permanent tokens, dynamic positions, current consultation state, emergency indicators, skipped-patient handling, call-next, start, skip, and completion operations.'),
        ('11.5 ETA and Live Patient Tracking', 'Patients can observe current serving state, dynamic queue position, and a bounded estimated waiting range. The ETA service considers queue order and consultation-time information rather than relying only on token arithmetic.'),
        ('11.6 Consultation and Prescription Workflow', 'Doctors can review the patient context, parse consultation text into diagnosis and medication fields, create prescriptions, attach supporting files, and complete the appointment workflow.'),
        ('11.7 Patient Memory and Contradiction Guard', 'The memory subsystem extracts structured facts such as allergies, conditions, medications, and procedures from source records. Facts retain provenance, are normalised and deduplicated, and may be marked conflicted when opposing assertions are detected.'),
        ('11.8 Records, Timelines, and Protected Uploads', 'Patients and authorised staff can view chronological medical activity, upload records, and protect selected files with password-based unlocking. The system hides protected URLs until a successful authorisation step.'),
        ('11.9 Hospital, Clinic, and Administration', 'Facility dashboards support roster and affiliation operations, while administrator dashboards expose statistics, audit logs, settings, and comparative benchmark results.'),
        ('11.10 Benchmark and Validation Tooling', 'Synthetic seeders, database inspection, acceptance tests, Kafka rebalance tests, memory tests, and scenario simulations support repeatable demonstrations and evaluation.'),
    ]
    for title, body in modules:
        add_heading(doc, title, 2)
        add_para(doc, body)

    add_heading(doc, '12. Clinical Queue and ETA Methodology', 1)
    add_heading(doc, '12.1 Permanent Token and Dynamic Position', 2)
    add_para(doc, 'LifeFile separates the permanent token number assigned to an appointment from the dynamic position at which the patient is likely to be served. This distinction preserves traceability: the appointment does not change identity merely because the service order changes. The queue can therefore respond to urgency while the patient and staff retain a stable reference.')
    add_heading(doc, '12.2 ACPA Clinical Priority Score', 2)
    add_para(doc, 'The ACPA implementation calculates a composite score from token order, triage level, waiting time, and missed calls. In conceptual form:')
    add_callout(doc, 'Composite priority model', 'CEP = token-order contribution + triage bonus + wait-time aging - bounded missed-call penalty', fill=LIGHT_BLUE)
    add_para(doc, 'Higher triage levels receive stronger urgency bonuses. Waiting time contributes an aging term that reduces the risk of routine patients being indefinitely displaced. Missed-call penalties are bounded so that a patient is deprioritised after being skipped but is not permanently excluded from service. The exact coefficients are implementation parameters and should be calibrated through clinical governance and observed operational data before real-world deployment.')
    add_heading(doc, '12.3 ETA Calculation', 2)
    add_para(doc, 'The ETA service derives a bounded range rather than a falsely precise point estimate. The estimate uses the patient\'s current dynamic position and consultation-time information. As queue events occur, the backend can recompute state and the frontend can receive an update through the real-time channel. A bounded range communicates uncertainty more responsibly than an exact promise, especially in clinical environments where consultation duration can vary materially.')
    add_heading(doc, '12.4 Check-In Crowd Control', 2)
    add_para(doc, 'The appointment workflow applies a time window around the scheduled slot. Before the opening threshold, the patient sees a locked state; during the active period, check-in is enabled; after the expiry threshold, the action is closed or requires an explicit staff workflow. This policy reduces premature crowding and provides a predictable rule that can be explained to patients.')

    add_heading(doc, '13. Patient Memory and Safety Workflow', 1)
    add_para(doc, 'The patient-memory layer is a derived context service. It does not replace source records. Instead, it creates reviewable memory items with category, type, content, confidence, source references, status, and conflict notes. This distinction is important because derived summaries may be incomplete, ambiguous, or wrong and must remain traceable to the material from which they were produced.')
    add_heading(doc, '13.1 Extraction Pipeline', 2)
    for item in ['Collect eligible source information from prescriptions, notes, diagnoses, medication fields, allergies, and procedures.', 'Apply deterministic extraction and normalisation rules for predictable patterns.', 'Optionally invoke Gemini-assisted structured extraction when configured.', 'Deduplicate related facts and associate each fact with one or more source records.', 'Compare new assertions with existing memory and identify possible contradictions.', 'Persist the result and make it available for patient, doctor, hospital, or administrator review according to role permissions.']:
        add_number(doc, item)
    add_heading(doc, '13.2 Contradiction Handling', 2)
    add_para(doc, 'A contradiction is treated as a safety signal, not as an automatic verdict. For example, a prior statement indicating no known drug allergy and a later statement identifying a penicillin allergy should be surfaced for review. The system can mark the memory item as conflicted, show the source context, and allow a correction or review workflow. A qualified professional remains responsible for deciding which statement is clinically current.')
    add_heading(doc, '13.3 Patient Correction and Review', 2)
    add_para(doc, 'Patients can request a correction to a derived memory item. Doctors or authorised reviewers can approve or reject the request, after which the memory record and audit history are updated. This creates a controlled feedback path and acknowledges that patients are important participants in the accuracy of their longitudinal context.')

    add_heading(doc, '14. Security, Privacy, and Governance', 1)
    add_para(doc, 'Because LifeFile handles sensitive healthcare information, security and governance are core design concerns rather than optional enhancements. The current prototype implements several useful controls, while also exposing areas that must be hardened before production use.')
    add_heading(doc, '14.1 Implemented Controls', 2)
    for item in ['JWT bearer authentication and protected frontend routes.', 'Role checks through middleware and role-specific workflows.', 'bcrypt-based password hashing for user credentials and protected medical-record access.', 'Authorisation logic for patient-memory access across patients, doctors, hospitals, and administrators.', 'Audit-log model and administrator-facing audit review.', 'Password-protected medical records whose file URLs are withheld until unlock succeeds.', 'Facility and hospital identifiers used to scope queues, rosters, and related data.']:
        add_bullet(doc, item)
    add_heading(doc, '14.2 Security Hardening Required for Production', 2)
    for item in ['Move browser token storage away from localStorage toward a hardened session strategy with appropriate XSS and CSRF protections.', 'Replace permissive CORS settings with an explicit allow-list and environment-specific configuration.', 'Protect uploaded files with authenticated, authorisation-aware delivery rather than relying on a broadly exposed static directory.', 'Complete the access-control TODO identified in the patient route and standardise authorisation checks across hospital, roster, and record endpoints.', 'Restrict or remove demonstration-only endpoints such as unauthenticated demo resynchronisation before deployment.', 'Make audit-log creation server-controlled and immutable from the perspective of ordinary users.', 'Make health reporting reflect actual MongoDB and Kafka availability rather than returning static flags.', 'Add rate limiting, input validation, structured logging, secret rotation, backup controls, retention policies, and security testing.', 'Conduct formal privacy, clinical safety, threat modelling, and regulatory review appropriate to the deployment jurisdiction.']:
        add_bullet(doc, item)
    add_callout(doc, 'Governance boundary', 'The platform should present triage, risk, memory, and queue outputs as decision support. It must not imply that a software score independently diagnoses a patient or determines treatment.', fill='FFF2CC')

    add_heading(doc, '15. Data Model and API Surface', 1)
    add_heading(doc, '15.1 Principal Data Models', 2)
    add_table(doc, ['Model', 'Purpose', 'Representative relationships'], [
        ('User', 'Identity, credentials, role, and account state', 'Links to patient, doctor, hospital, or administrative profile'),
        ('Patient', 'Patient demographic and clinical profile', 'Appointments, prescriptions, medical records, memories'),
        ('Doctor', 'Professional profile, schedule, affiliations', 'Hospital, clinic, appointments, reviews'),
        ('Hospital / Clinic', 'Facility identity and operations', 'Doctors, rosters, appointments, hospital analytics'),
        ('Appointment', 'Booking, token, status, triage, timing, and queue metadata', 'Patient, doctor, hospital, prescription'),
        ('Prescription / MedicalRecord', 'Authoritative clinical outputs and uploaded evidence', 'Patient, appointment, source for memory'),
        ('PatientMemory', 'Derived structured context with provenance and conflict state', 'Patient, source records, corrections, audits'),
        ('MemoryCorrection', 'Patient-requested change and reviewer decision', 'PatientMemory, patient, reviewer, audit'),
        ('AuditLog', 'Traceability for sensitive or administrative actions', 'Actor, resource, action, timestamp'),
        ('Review / JoinRequest', 'Feedback and facility affiliation workflows', 'Patient-doctor or doctor-hospital relationships'),
    ], widths=[1.45, 3.0, 2.15])
    add_heading(doc, '15.2 Route Families', 2)
    add_table(doc, ['Route family', 'Representative capabilities'], [
        ('/api/auth', 'Registration, login, current user'),
        ('/api/appointments', 'Booking, rescheduling, cancellation, walk-ins, check-in, status'),
        ('/api/queue', 'Queue retrieval, ETA, call-next, skip, completion, resynchronisation'),
        ('/api/memory', 'Memory retrieval, extraction, sources, corrections, review'),
        ('/api/patients and /api/prescriptions', 'Patient profiles, records, uploads, prescriptions, attachments'),
        ('/api/doctors, /api/hospitals, /api/clinics', 'Profiles, facilities, rosters, affiliations, administration'),
        ('/api/analytics, /api/stats, /api/benchmark', 'Operational analytics, statistics, simulations, comparisons'),
        ('/api/audit-logs and /api/reviews', 'Audit access and feedback workflows'),
    ], widths=[2.0, 4.6])

    add_heading(doc, '16. Technology Stack and Tools', 1)
    add_table(doc, ['Category', 'Selected technology', 'Reason for selection'], [
        ('Frontend', 'React 18, TypeScript, Vite', 'Component-based UI, type-aware development, fast local iteration'),
        ('Design and UI', 'Tailwind CSS, Lucide React, Recharts', 'Consistent styling, accessible iconography, operational dashboards and charts'),
        ('Client data', 'Axios, TanStack React Query, Zustand', 'HTTP communication, server-state caching, lightweight application state'),
        ('Backend', 'Node.js, Express.js', 'Mature API ecosystem and straightforward modular route composition'),
        ('Persistence', 'MongoDB, Mongoose', 'Flexible document modelling for healthcare operational entities and derived memory'),
        ('Authentication', 'JWT, bcryptjs', 'Token-based API access and password hashing'),
        ('Real time', 'Apache Kafka, KafkaJS, Socket.IO', 'Event transport plus browser-facing live updates'),
        ('NLP and AI assistance', 'compromise, deterministic rules, optional Gemini integration', 'Local explainable heuristics with optional structured extraction assistance'),
        ('File handling', 'Multer and local upload storage', 'Prototype support for signatures and medical attachments'),
        ('Validation', 'Node-based acceptance tests, benchmark simulations, memory and Kafka tests', 'Behavioural verification of high-risk operational paths'),
    ], widths=[1.25, 2.2, 3.15])

    add_heading(doc, '17. Development Methodology', 1)
    add_para(doc, 'The recommended development methodology is iterative and evidence-driven. Each workflow should be implemented from the domain model outward: define the state transition, enforce it at the backend boundary, expose it through the frontend, publish the relevant event, and verify the resulting behaviour through an executable scenario.')
    phases = [
        ('Phase 1', 'Requirements, domain vocabulary, risk boundaries, and stakeholder workflows', '1-2 weeks'),
        ('Phase 2', 'Architecture, data model, route contracts, authentication baseline, local environment', '1-2 weeks'),
        ('Phase 3', 'Patient discovery, doctor profiles, booking, appointment state, and check-in', '2-3 weeks'),
        ('Phase 4', 'ACPA queue, ETA service, consultation states, and real-time event delivery', '3-4 weeks'),
        ('Phase 5', 'Records, prescriptions, memory extraction, corrections, and audit visibility', '3-4 weeks'),
        ('Phase 6', 'Hospital/admin workflows, analytics, benchmarks, seed data, and demo preparation', '2-3 weeks'),
        ('Phase 7', 'Security hardening, acceptance testing, documentation, and final evaluation', '2-3 weeks'),
    ]
    add_table(doc, ['Phase', 'Major activities', 'Indicative duration'], phases, widths=[1.0, 4.45, 1.15])
    add_heading(doc, '17.1 Engineering Practices', 2)
    for item in ['Use small, explicit route modules and service boundaries.', 'Treat server-side state transitions as authoritative and idempotent where possible.', 'Keep derived memory separate from source clinical records and preserve provenance.', 'Use synthetic seed data for demonstrations and never expose real patient data in presentation environments.', 'Record assumptions, limitations, and test evidence alongside feature claims.', 'Prefer focused acceptance tests for queue, memory, authentication, authorisation, and event-recovery paths.']:
        add_bullet(doc, item)

    add_heading(doc, '18. Validation and Testing Strategy', 1)
    add_para(doc, 'LifeFile includes executable validation assets rather than relying only on screenshots or manual demonstrations. The validation strategy should be understood as prototype-level evidence: it demonstrates intended behaviour and mathematical properties, but it does not by itself establish clinical effectiveness, production resilience, or regulatory compliance.')
    add_table(doc, ['Validation area', 'Available evidence / command', 'What it should demonstrate'], [
        ('Memory extraction', 'test-memory-engine.js', 'Deterministic extraction, normalisation, contradiction handling, and derived-memory behaviour'),
        ('Memory authorisation', 'test-memory-auth.js', 'Role-aware access and correction/review permissions'),
        ('Queue acceptance', 'test-acceptance-queue.js', 'Check-in, dynamic ordering, call-next, skip, and completion flows'),
        ('Kafka resilience', 'test-kafka-rebalance.js', 'Consumer/rebalance behaviour and event-handling safeguards'),
        ('Benchmark suite', 'sih-validation/cli.js and sih-validation/tests/', 'Repeatable comparison of FIFO, priority, ACPA, and ablation scenarios'),
        ('Database inspection', 'inspect-db.js', 'Human-readable inspection of users, appointments, memories, and audit state'),
        ('Frontend build', 'npm run build in scos-frontend', 'TypeScript compilation and Vite production build'),
    ], widths=[1.35, 2.4, 2.85])
    add_heading(doc, '18.1 Demonstration Acceptance Scenarios', 2)
    for item in ['A high-severity complaint receives a higher assistive triage level and persists with the appointment.', 'A patient outside the configured check-in window cannot join prematurely.', 'A dynamic queue promotes an urgent checked-in patient without changing the permanent token.', 'A doctor state transition becomes visible to a second client through the real-time path.', 'A conflicting allergy statement appears as a reviewable warning with provenance.', 'A patient correction request follows a controlled approval or rejection path.', 'A facility-scoped user sees only the queue and records permitted by the active facility context.', 'A benchmark run produces reproducible results for the selected seed and scenario.']:
        add_bullet(doc, item)

    add_heading(doc, '19. Feasibility Study', 1)
    add_heading(doc, '19.1 Technical Feasibility', 2)
    add_para(doc, 'The prototype uses mature and accessible technologies. React, Express, MongoDB, KafkaJS, Socket.IO, and TypeScript are widely supported, and the project already contains route modules, services, models, seeders, and validation scripts. The main technical challenge for future work is not basic implementation but production hardening: reliable deployment, privacy controls, observability, concurrency guarantees, and clinical validation.')
    add_heading(doc, '19.2 Economic Feasibility', 2)
    add_para(doc, 'A demonstration installation can be developed using open-source software and standard development hardware. Costs may arise from managed MongoDB or Kafka hosting, file storage, observability, domain and deployment services, and optional generative-AI usage. A staged architecture allows the project to operate locally without requiring specialised hardware.')
    add_heading(doc, '19.3 Operational Feasibility', 2)
    add_para(doc, 'The role-specific design maps naturally to existing clinic responsibilities. Patients receive a simpler, status-oriented workflow; doctors receive operational controls; facility managers receive scoped oversight; and administrators receive network-level visibility. Adoption would still require training, process alignment, role governance, support procedures, and an explicit policy for handling automated recommendations.')
    add_heading(doc, '19.4 Legal and Ethical Feasibility', 2)
    add_para(doc, 'The project is feasible as an academic prototype when it uses synthetic data, states its clinical limitations, and treats AI outputs as assistive. A real deployment would require jurisdiction-specific privacy and health-data review, consent and retention policies, access logging, incident response, clinical safety assessment, vendor review, and formal governance.')

    add_heading(doc, '20. Expected Outcomes', 1)
    for item in ['A working multi-role healthcare operations prototype accessible through a web interface.', 'A dynamic queue demonstration showing the interaction of urgency, fairness, missed calls, and time-aware check-in.', 'A real-time state synchronisation path connecting backend events to patient and doctor views.', 'A patient-memory subsystem with source links, contradiction warnings, correction requests, and review states.', 'A coherent data model and REST API surface covering users, facilities, appointments, records, prescriptions, memories, analytics, and audits.', 'A repeatable validation and benchmark package suitable for academic demonstration and further experimentation.', 'A professional technical foundation that can be extended toward production hardening and external integration.']:
        add_bullet(doc, item)

    add_heading(doc, '21. Innovation and Novelty', 1)
    add_number(doc, 'Operational integration: brings booking, check-in, queueing, consultation, records, memory, and governance into one workflow.')
    add_number(doc, 'Ethically framed dynamic priority: represents urgency without discarding fairness or traceability.')
    add_number(doc, 'Derived context with provenance: extracts useful patient context while preserving links to authoritative source records.')
    add_number(doc, 'Real-time transparency: makes queue and consultation state visible to patients and staff through event-driven updates.')
    add_number(doc, 'Human-reviewable automation: uses NLP and optional generative assistance as an aid, with correction and review paths rather than opaque final decisions.')
    add_number(doc, 'Benchmark-oriented evaluation: treats queue policy as an experiment that can be compared through scenarios rather than only presented as a UI feature.')

    add_heading(doc, '22. Limitations and Responsible-Use Boundaries', 1)
    add_para(doc, 'The following limitations are material to an honest evaluation of the current implementation:')
    for item in ['Patient-facing symptom triage and risk estimation are heuristic or locally rule-based and should not be interpreted as clinical diagnosis.', 'Generative-AI extraction may be unavailable, variable, or incorrect; deterministic fallback and human review remain important.', 'The current prototype uses local upload storage and development-oriented configuration that require hardening for production.', 'The frontend authentication token strategy, permissive CORS configuration, static upload exposure, and identified access-control TODOs require remediation.', 'The validation assets are targeted executable tests and simulations; there is no stated root-level CI pipeline or comprehensive frontend test suite in the package manifests.', 'Synthetic demo accounts and data are appropriate for presentations but must never be confused with real patient records.', 'Queue coefficients, ETA assumptions, and clinical escalation policies require governance and calibration with domain experts.']:
        add_bullet(doc, item)
    add_callout(doc, 'Responsible-use statement', 'LifeFile is a coordination and decision-support prototype. It does not independently diagnose, prescribe, determine a patient\'s eligibility for care, or replace a qualified healthcare professional.', fill='FFF2CC')

    add_heading(doc, '23. Project Timeline', 1)
    add_table(doc, ['Milestone', 'Deliverable', 'Success indicator'], [
        ('M1', 'Requirements and architecture baseline', 'Approved workflows, data dictionary, risk register, and component map'),
        ('M2', 'Core identity and appointment workflows', 'Role login, booking, appointment state, and check-in operate end to end'),
        ('M3', 'Queue and real-time integration', 'ACPA ordering, ETA, consultation state, and client synchronisation are demonstrable'),
        ('M4', 'Clinical context and record workflows', 'Prescriptions, records, memory extraction, conflicts, and corrections are traceable'),
        ('M5', 'Administration and evaluation tooling', 'Facility controls, audit visibility, analytics, seeders, and benchmarks are usable'),
        ('M6', 'Hardening and final presentation', 'Focused tests pass, limitations are documented, and the live demonstration is repeatable'),
    ], widths=[0.75, 2.55, 3.8])

    add_heading(doc, '24. Team Members and Responsibilities', 1)
    add_para(doc, 'The following table is intentionally left ready for institutional completion. Names, roll numbers, guide details, and individual ownership should be filled from the official project records rather than inferred from the repository.')
    add_table(doc, ['Sr. No.', 'Name', 'Roll No.', 'Role / responsibility'], [
        ('1', '[Team member name]', '[Roll no.]', 'Project coordination, architecture, and integration'),
        ('2', '[Team member name]', '[Roll no.]', 'Frontend experience, UI workflows, and accessibility'),
        ('3', '[Team member name]', '[Roll no.]', 'Backend APIs, data models, and authentication'),
        ('4', '[Team member name]', '[Roll no.]', 'Algorithms, memory services, testing, and evaluation'),
    ], widths=[0.7, 1.75, 1.0, 3.15])
    add_heading(doc, '24.1 Project Guide', 2)
    add_table(doc, ['Field', 'Value'], [
        ('Guide name', '[Faculty member name]'),
        ('Department', 'Department of Information Technology'),
        ('Institution', '[Institution name]'),
        ('Academic year', '2026-2027'),
    ], widths=[2.0, 4.6])

    add_heading(doc, '25. References', 1)
    refs = [
        '[1] LifeFile project repository and local implementation, github.com/lavyadamania/LifeFile.',
        '[2] LifeFile / SCOS System Documentation, project_overview.md.',
        '[3] LifeFile Feature Inventory, LIFEFILE_FEATURE_INVENTORY.md.',
        '[4] LifeFile Presentation Flow and Demonstration Manual, PRESENTATION_FLOW.md.',
        '[5] React documentation, https://react.dev/.',
        '[6] Express documentation, https://expressjs.com/.',
        '[7] MongoDB documentation, https://www.mongodb.com/docs/.',
        '[8] Apache Kafka documentation, https://kafka.apache.org/documentation/.',
        '[9] Socket.IO documentation, https://socket.io/docs/v4/.',
        '[10] JSON Web Token specification and documentation, https://jwt.io/introduction.',
        '[11] NIST, Privacy Framework: A Tool for Improving Privacy through Enterprise Risk Management, National Institute of Standards and Technology.',
        '[12] World Health Organization, Ethics and Governance of Artificial Intelligence for Health, WHO guidance document.',
    ]
    for ref in refs:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(ref)

    add_heading(doc, '26. Conclusion', 1)
    add_para(doc, 'LifeFile represents a substantial full-stack effort to make healthcare operations more coordinated, transparent, and context-aware. Its strongest contribution is the integration of workflows that are often treated as separate concerns: appointments, time-aware check-in, dynamic queue management, live status, clinical documentation, patient memory, facility administration, and evaluation.')
    add_para(doc, 'The ACPA queue engine provides the project with a clear operational centre: it demonstrates how a clinic can respond to urgency while retaining stable identifiers and fairness-oriented aging. The patient-memory subsystem provides a second differentiator by converting scattered clinical information into reviewable context with provenance and contradiction handling. Kafka and Socket.IO demonstrate the event-driven path required to keep multiple views aligned as care progresses.')
    add_para(doc, 'The project is best evaluated as an intelligent healthcare operations prototype with a credible foundation for future work. Its next level of maturity is defined by disciplined hardening: stronger authorisation, safer file delivery, explicit privacy governance, observable infrastructure, clinical review, reproducible deployment, and broader automated testing. With those additions, LifeFile can progress from a high-value academic demonstration toward a more robust platform for controlled research and operational experimentation.')
    add_callout(doc, 'Final positioning', 'LifeFile does not attempt to automate away clinical judgement. It organises the information, state, and communication around care so that patients and professionals can act with better visibility and stronger operational support.')

    doc.core_properties.title = 'LifeFile Unified Project Proposal'
    doc.core_properties.subject = 'Professional project proposal and technical documentation for the LifeFile Smart Clinic Operating System'
    doc.core_properties.author = 'LifeFile Project Team'
    doc.core_properties.keywords = 'LifeFile, SCOS, healthcare operations, dynamic queue, patient memory, real-time systems'
    doc.save(OUTPUT)


if __name__ == '__main__':
    build()
    print(OUTPUT)
